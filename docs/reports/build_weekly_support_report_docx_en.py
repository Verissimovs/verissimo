from pathlib import Path
from zipfile import ZipFile

from docx import Document

from build_relatorio_apoio_semanal_docx import (
    LIGHT_BLUE,
    LIGHT_RED,
    LIGHT_YELLOW,
    add_bullets,
    add_callout,
    add_footer,
    add_heading,
    add_label_value,
    add_numbered,
    add_paragraph,
    add_table,
    configure_document,
)


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "weekly-support-report-verissimo.docx"


WEEKLY_FEATURES = [
    {
        "week": "Week 1",
        "focus": "Technical foundation and deterministic rules engine",
        "features": [
            "Next.js project scaffolded with TypeScript, Vitest, ESLint and production build support.",
            "Zod schema created for declarative benefit rule files.",
            "Deterministic operators implemented: gte, lte, eq, in and is_true.",
            "Outcome model implemented with four states: Worth exploring, May be relevant, Need more answers and Not a match.",
            "Rule loader added with validation, referential integrity checks and rulesVersion hashing.",
            "Initial Attendance Allowance rule file created.",
        ],
        "evidence": "Rules-engine unit tests, coverage above the defined gate and a working Next.js production build.",
        "report_link": "Supports Methodology, Design and Implementation, and Testing chapters.",
    },
    {
        "week": "Week 2",
        "focus": "Benefit expansion and evidence recommendations",
        "features": [
            "Universal Credit and Personal Independence Payment added as YAML rule files.",
            "assessAll implemented to evaluate all three benefits together.",
            "Deterministic evidence recommendations added based on triggered criteria.",
            "Tests strengthened to prevent hard-coded criteria or invented source URLs.",
        ],
        "evidence": "Local test suite, coverage gate and versioned YAML files.",
        "report_link": "Supports Requirements, Implementation and Verification sections.",
    },
    {
        "week": "Week 3",
        "focus": "Guided questionnaire and assessment results",
        "features": [
            "Wizard-style intake interface implemented with one question shown at a time.",
            "Question configuration mapped to every answer key required by the rule files.",
            "Internal POST /api/assess route added so YAML loading and file-system access stay server-side.",
            "Result cards added with summary, decisive criterion, evidence recommendations and expandable rule trace.",
            "Initial public deployment completed on Vercel.",
        ],
        "evidence": "Vitest, TypeScript, ESLint, build verification, local smoke test and production smoke test.",
        "report_link": "Supports User Interface Design, System Architecture and Deployment sections.",
    },
    {
        "week": "Week 4",
        "focus": "Explanation layer and bounded LLM integration",
        "features": [
            "Deterministic template explanation layer implemented.",
            "Internal POST /api/explain route added.",
            "Prompt builder and validator added to constrain LLM outputs.",
            "Gemini integrated through a server-side API key.",
            "Deterministic fallback preserved for missing API keys, provider errors or unsafe model output.",
        ],
        "evidence": "Template, prompt, validation, route and Gemini adapter tests.",
        "report_link": "Supports AI Integration, Safety, Explainability and Limitations sections.",
    },
    {
        "week": "Week 5",
        "focus": "Printable report, accessibility and evaluation readiness",
        "features": [
            "Per-benefit printable report implemented with metadata, rule trace, evidence and sources.",
            "Static accessibility review completed and ARIA issues fixed.",
            "Playwright and axe-core added for automated WCAG-oriented checks.",
            "Keyboard-only navigation test added.",
            "Readability audit added using Flesch Reading Ease and Flesch-Kincaid checks.",
            "Ethics and usability-testing pack prepared.",
        ],
        "evidence": "npm run test:a11y, keyboard test, readability test and docs/ethics materials.",
        "report_link": "Supports Ethics, Accessibility, Evaluation Plan and Results sections.",
    },
    {
        "week": "Week 6 / current refinement",
        "focus": "CPAG source strategy, modern UI and demonstration readiness",
        "features": [
            "CPAG Welfare Rights made the primary source family shown in rule trace.",
            "GOV.UK retained as complementary official guidance.",
            "Blue Decision Lab visual remake implemented with a modern technology-oriented interface.",
            "Result cards polished with shorter badges, cleaner surfaces and lighter repeated actions.",
            "Anonymous study JSON export added for usability evaluation.",
            "Production deployment updated at https://verissimo-tawny.vercel.app.",
        ],
        "evidence": "119 tests passing, TypeScript, lint, coverage, build, accessibility tests and production smoke test.",
        "report_link": "Supports Final Artefact, Evaluation Readiness and Demonstration sections.",
    },
    {
        "week": "Week 7 / next step",
        "focus": "Participant testing and feedback collection",
        "features": [
            "Submit the ethics pack and wait for approval before involving real participants.",
            "Run usability testing with 5 to 8 participants using fictional scenarios.",
            "Collect anonymous JSON exports and post-test questionnaire responses.",
            "Analyse comprehension, trust, perceived accessibility and points of confusion.",
        ],
        "evidence": "Consent form, participant information sheet, screening form, task list, script and questionnaire.",
        "report_link": "Will support Methodology, Evaluation, Results and Further Work chapters.",
    },
]


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_executive_summary(doc)
    add_weekly_feature_plan(doc)
    add_architecture_summary(doc)
    add_data_sources_summary(doc)
    add_testing_summary(doc)
    add_ethics_summary(doc)
    add_report_writing_map(doc)
    add_risks_and_next_steps(doc)
    add_appendix_evidence(doc)
    add_footer(doc)
    doc.save(OUTPUT)
    structural_audit(OUTPUT)
    print(OUTPUT)


def add_cover(doc: Document) -> None:
    add_heading(doc, "VERISSIMO Weekly Support Report", 1)
    add_paragraph(
        doc,
        "A weekly presentation and dissertation-writing support document for the VERISSIMO benefits decision-support prototype.",
    )
    rows = [
        ["Field", "Value"],
        ["Project", "VERISSIMO"],
        ["Student", "Viviane Verissimo dos Santos"],
        ["Course", "BSc (Hons) Computing - Southampton Solent University"],
        ["Date", "20/07/2026"],
        ["Public link", "https://verissimo-tawny.vercel.app"],
        ["Document purpose", "Support weekly presentations and provide structured material for the academic report."],
    ]
    add_table(doc, rows, [2200, 7160])
    add_callout(
        doc,
        "Core message",
        "VERISSIMO separates decision-making from explanation. A deterministic rules engine calculates assessment results, while AI is used only to rewrite explanations in simpler language. This makes the prototype more auditable, safer and more suitable for academic evaluation.",
        LIGHT_BLUE,
    )


def add_executive_summary(doc: Document) -> None:
    add_heading(doc, "1. Executive Summary", 1)
    add_paragraph(
        doc,
        "VERISSIMO is a web prototype designed to help users review possible UK social-security benefits. The application guides users through simple questions, evaluates answers against declarative rules and presents results with traceable sources, recommended evidence and simplified explanations.",
    )
    add_paragraph(
        doc,
        "The current version covers Attendance Allowance, Universal Credit and Personal Independence Payment. The project prioritises auditability, privacy, accessibility and clarity. It does not replace GOV.UK, CPAG, the Department for Work and Pensions, welfare advisers or official decisions.",
    )
    add_bullets(
        doc,
        [
            "Main artefact: a Next.js application deployed publicly on Vercel.",
            "Decision method: deterministic YAML-based rules engine.",
            "Primary rule-trace source: CPAG Welfare Rights.",
            "Complementary source: GOV.UK.",
            "AI role: optional Gemini explanation only; it does not decide eligibility.",
            "Current state: 119 local tests passing and a working public deployment.",
        ],
    )


def add_weekly_feature_plan(doc: Document) -> None:
    add_heading(doc, "2. Weekly Feature Breakdown", 1)
    add_paragraph(
        doc,
        "This section organises the project into weekly blocks so progress can be presented clearly in supervision meetings and reused in the academic report. Each week lists the focus, delivered features, evidence and report-writing value.",
    )
    for item in WEEKLY_FEATURES:
        add_heading(doc, f"{item['week']} - {item['focus']}", 2)
        add_bullets(doc, item["features"])
        add_label_value(doc, "Evidence", item["evidence"])
        add_label_value(doc, "Report-writing use", item["report_link"])

    rows = [["Week", "Focus", "Main deliverables", "Evidence"]]
    for item in WEEKLY_FEATURES:
        rows.append([item["week"], item["focus"], "; ".join(item["features"][:3]), item["evidence"]])
    add_heading(doc, "Weekly Summary Table", 2)
    add_table(doc, rows, [1150, 1900, 4010, 2300])


def add_architecture_summary(doc: Document) -> None:
    add_heading(doc, "3. Application Architecture", 1)
    add_paragraph(
        doc,
        "The architecture keeps eligibility decisions outside the LLM. The frontend presents the questionnaire and results. The internal API receives answers and calls the server-side rules engine. YAML files store declarative benefit criteria.",
    )
    rows = [
        ["Layer", "Responsibility", "Main files"],
        ["Frontend", "Guided questionnaire, result cards, printable report and anonymous export.", "src/components/assessment/AssessmentApp.tsx; src/app/globals.css"],
        ["API", "Internal assessment and explanation routes.", "src/app/api/assess/route.ts; src/app/api/explain/route.ts"],
        ["Rules engine", "Load YAML, validate schema, evaluate criteria, derive outcome and recommend evidence.", "src/lib/rules/engine.ts; load.ts; schema.ts; outcome.ts; operators.ts"],
        ["Rules", "Declarative criteria, sources and evidence by benefit.", "rules/attendance-allowance.yaml; universal-credit.yaml; personal-independence-payment.yaml"],
        ["Explanation", "Deterministic template, safe prompt, validator and Gemini adapter.", "src/lib/explanations/template.ts; llm.ts; gemini.ts"],
        ["Tests", "Unit, route, UI, accessibility, keyboard and readability tests.", "tests/rules; tests/intake; tests/explanations; tests/e2e"],
    ]
    add_table(doc, rows, [1450, 3600, 4310])
    add_callout(
        doc,
        "Important architectural decision",
        "The LLM does not calculate eligibility, create criteria or choose benefits. It receives a result already produced by the rules engine and attempts to explain it in simpler language. If the external response is unsafe or unavailable, the application uses a deterministic fallback.",
        LIGHT_YELLOW,
    )


def add_data_sources_summary(doc: Document) -> None:
    add_heading(doc, "4. Data Sources and Traceability", 1)
    add_paragraph(
        doc,
        "CPAG Welfare Rights is the primary source family shown in the rule trace. GOV.UK remains available as complementary official guidance. The application does not scrape sources live: rules are explicitly modelled in YAML so behaviour remains reproducible and testable.",
    )
    rows = [
        ["Benefit", "Primary CPAG source", "Complementary GOV.UK source"],
        ["Universal Credit", "https://cpag.org.uk/welfare-rights/key-topics/universal-credit/universal-credit-basics", "https://www.gov.uk/universal-credit/eligibility"],
        ["Personal Independence Payment", "https://cpag.org.uk/welfare-rights/key-topics/personal-independence-payment", "https://www.gov.uk/pip/eligibility"],
        ["Attendance Allowance", "https://cpag.org.uk/welfare-rights", "https://www.gov.uk/attendance-allowance/eligibility"],
    ]
    add_table(doc, rows, [1800, 4180, 3380])
    add_bullets(
        doc,
        [
            "Each result includes rulesVersion so the exact YAML version can be traced.",
            "Each criterion includes an explanation and a source URL.",
            "Complementary GOV.UK sources are shown separately from the CPAG rule trace.",
            "Avoiding live scraping reduces instability during assessment and demonstration.",
        ],
    )


def add_testing_summary(doc: Document) -> None:
    add_heading(doc, "5. Quality Evidence", 1)
    add_paragraph(
        doc,
        "The project includes automated checks that can be cited as engineering, reliability and accessibility evidence. The latest verification after result-card polishing passed the following gates.",
    )
    rows = [
        ["Check", "Recent result", "What it supports"],
        ["npm test", "14 files / 119 tests passing", "Rules, routes, UI, explanations, helpers and readability."],
        ["npx tsc --noEmit", "Passed", "TypeScript type consistency."],
        ["npm run lint", "Passed", "Static code-quality baseline."],
        ["npm run test:coverage", "97.43% statements; 94.2% branches", "Coverage above the defined gate."],
        ["npm run build", "Passed", "Next.js production build."],
        ["npm run test:a11y", "3 Playwright tests passing", "Automated axe/WCAG checks and keyboard-only flow."],
        ["Production smoke test", "HTTP 200 and API returns CPAG + GOV.UK", "Public deployment works after changes."],
    ]
    add_table(doc, rows, [2300, 2700, 4360])


def add_ethics_summary(doc: Document) -> None:
    add_heading(doc, "6. Ethics, Privacy and User Evaluation", 1)
    add_paragraph(
        doc,
        "The project is structured to reduce risk before any participant testing. There is no login, user database or persistent storage of personal data. User testing should use fictional scenarios until formal ethics approval is granted.",
    )
    add_bullets(
        doc,
        [
            "Participants should not enter real health, income, immigration or benefits information.",
            "The study JSON export is anonymous and intended for usability evaluation, not profiling.",
            "The site states that it does not provide official advice.",
            "The docs/ethics folder contains the consent form, participant information sheet, screening form, task list, script and post-test questionnaire.",
        ],
    )
    add_callout(
        doc,
        "Weekly meeting point",
        "Before real usability testing, the next key milestone is to review and submit the ethics pack through the university portal.",
        LIGHT_RED,
    )


def add_report_writing_map(doc: Document) -> None:
    add_heading(doc, "7. How This Supports the Academic Report", 1)
    rows = [
        ["Report section", "Material to use", "Local evidence"],
        ["Introduction / Problem", "Complexity of benefits, risk of AI-led decisions and need for explainability.", "Executive summary and prototype limitations."],
        ["Requirements", "Three benefits, traceable sources, accessibility, privacy, simple explanations and public deployment.", "Specs in docs/specs."],
        ["Methodology", "Incremental development, TDD, continuous verification, UI reviews and usability-testing plan.", "WORKLOG.md and automated tests."],
        ["Design and Implementation", "Next.js architecture, internal API, YAML rules engine, CPAG/GOV.UK strategy, LLM boundary and Blue Decision Lab UI.", "src/, rules/, docs/plans and screenshots."],
        ["Testing", "Vitest, coverage, TypeScript, ESLint, Playwright, axe-core, keyboard-only flow and readability audit.", "tests/ and coverage/."],
        ["Evaluation", "Plan for 5 to 8 participants using fictional scenarios and anonymous JSON export.", "docs/ethics and Export study JSON."],
        ["Conclusion / Further Work", "More benefits, rule-editing interface, adviser review, source-update workflow and longer-term evaluation.", "Risks and next steps in this document."],
    ]
    add_table(doc, rows, [1800, 4470, 3090])


def add_risks_and_next_steps(doc: Document) -> None:
    add_heading(doc, "8. Risks, Limitations and Next Steps", 1)
    rows = [
        ["Level", "Risk or limitation", "Proposed mitigation"],
        ["High", "Benefit rules can change and the prototype must not provide official advice.", "Track last_verified, show sources and direct users to CPAG, GOV.UK or a qualified adviser."],
        ["High", "Real-user testing before ethics approval.", "Use fictional scenarios only until formal approval is granted."],
        ["Medium", "The LLM may produce unsuitable explanatory text.", "Validate responses, allowlist URLs and retain deterministic fallback."],
        ["Medium", "The prototype covers only three benefits.", "Present it as a focused prototype and list expansion as further work."],
        ["Low", "The UI may need adjustment after participant feedback.", "Use JSON export, post-test questionnaire and documented visual iterations."],
    ]
    add_table(doc, rows, [900, 4300, 4160])
    add_heading(doc, "Recommended Next Steps", 2)
    add_numbered(
        doc,
        [
            "Review this weekly report with the supervisor or stakeholder.",
            "Finalise and submit the ethics application.",
            "Prepare the public-site demonstration script.",
            "Run usability testing with 5 to 8 participants after approval.",
            "Turn the sections in this document into academic report chapters.",
        ],
    )


def add_appendix_evidence(doc: Document) -> None:
    add_heading(doc, "9. Main Evidence and Artefacts", 1)
    add_bullets(
        doc,
        [
            "Public application: https://verissimo-tawny.vercel.app",
            "Rules: rules/attendance-allowance.yaml, rules/universal-credit.yaml, rules/personal-independence-payment.yaml",
            "Rules engine: src/lib/rules/",
            "UI: src/components/assessment/AssessmentApp.tsx and src/app/globals.css",
            "LLM boundary and explanation layer: src/lib/explanations/",
            "Tests: tests/",
            "Ethics pack: docs/ethics/",
            "Recent visual evidence: test-results/result-card-polish-desktop.png",
            "Latest cited deployment: dpl_EsNxnGJ5kKT31J4hwStxprmnyJdt",
        ],
    )


def structural_audit(path: Path) -> None:
    with ZipFile(path) as docx:
        document_xml = docx.read("word/document.xml").decode("utf-8")
        styles_xml = docx.read("word/styles.xml").decode("utf-8")
    required = [
        "VERISSIMO Weekly Support Report",
        "Weekly Feature Breakdown",
        "Week 1",
        "Week 6 / current refinement",
        "CPAG Welfare Rights",
        "119 local tests passing",
        "How This Supports the Academic Report",
    ]
    missing = [text for text in required if text not in document_xml]
    if missing:
        raise RuntimeError(f"Missing required document text: {missing}")
    forbidden = ["TODO", "TBD", "PLACEHOLDER"]
    found = [text for text in forbidden if text in document_xml]
    if found:
        raise RuntimeError(f"Forbidden placeholders found: {found}")
    if "Heading1" not in styles_xml and "Heading 1" not in styles_xml:
        raise RuntimeError("Heading styles not present in styles.xml")


if __name__ == "__main__":
    main()

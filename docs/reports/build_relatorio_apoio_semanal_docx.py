from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "relatorio-apoio-semanal-verissimo.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
INK = RGBColor(20, 32, 47)
MUTED = RGBColor(85, 98, 115)
WHITE = RGBColor(255, 255, 255)
LIGHT_BLUE = "E8F1FB"
LIGHT_GRAY = "F2F4F7"
LIGHT_YELLOW = "FFF7E6"
LIGHT_RED = "FCECEC"
LIGHT_GREEN = "EAF6EF"


WEEKLY_FEATURES = [
    {
        "week": "Semana 1",
        "focus": "Fundacao tecnica e motor de regras",
        "features": [
            "Projeto Next.js com TypeScript, Vitest, ESLint e build de producao.",
            "Schema Zod para arquivos YAML de beneficios.",
            "Operadores deterministicos: gte, lte, eq, in e is_true.",
            "Motor de outcome com quatro estados: Worth exploring, May be relevant, Need more answers e Not a match.",
            "Loader de regras com validacao, integridade referencial e rulesVersion por hash.",
            "Primeira versao do Attendance Allowance.",
        ],
        "evidence": "Testes unitarios do motor de regras, cobertura acima de 80%, build Next.js funcional.",
        "report_link": "Pode alimentar Methodology, Design and Implementation, e Testing.",
    },
    {
        "week": "Semana 2",
        "focus": "Expansao de beneficios e evidencias",
        "features": [
            "Universal Credit e Personal Independence Payment adicionados aos YAMLs.",
            "Avaliacao simultanea dos tres beneficios no assessAll.",
            "Recomendacao deterministica de evidencias conforme criterios acionados.",
            "Fortalecimento dos testes para impedir criterios ou URLs inventados.",
        ],
        "evidence": "Suite local, coverage gate e arquivos YAML versionados.",
        "report_link": "Pode alimentar Requirements, Implementation e Verification.",
    },
    {
        "week": "Semana 3",
        "focus": "Questionario guiado e resultados",
        "features": [
            "Interface de intake em formato wizard, uma pergunta por vez.",
            "Perguntas mapeadas para as chaves exigidas pelas regras.",
            "API interna POST /api/assess mantendo YAML e fs no servidor.",
            "Cards de resultado com resumo, criterio decisivo, evidencias e trace expansivel.",
            "Deploy inicial na Vercel com alias publico.",
        ],
        "evidence": "Vitest, TypeScript, ESLint, build, smoke local e smoke de producao.",
        "report_link": "Pode alimentar User Interface Design, System Architecture e Deployment.",
    },
    {
        "week": "Semana 4",
        "focus": "Camada de explicacao e LLM controlado",
        "features": [
            "Explicacao deterministica por template.",
            "Endpoint POST /api/explain.",
            "Prompt builder e validador para limitar o que o LLM pode dizer.",
            "Gemini conectado via API key server-side.",
            "Fallback deterministico quando nao ha chave, erro externo ou resposta insegura.",
        ],
        "evidence": "Testes de template, prompt, validacao, rota e adaptador Gemini.",
        "report_link": "Pode alimentar AI Integration, Safety, Explainability e Limitations.",
    },
    {
        "week": "Semana 5",
        "focus": "Relatorio imprimivel, acessibilidade e avaliacao",
        "features": [
            "Relatorio imprimivel por beneficio com metadata, regras, fontes e evidencias.",
            "Auditoria estatica de acessibilidade e ajustes de ARIA.",
            "Playwright com axe-core para verificacao WCAG automatizada.",
            "Teste de navegacao apenas por teclado.",
            "Auditoria automatizada de legibilidade com Flesch Reading Ease e Flesch-Kincaid.",
            "Pacote de etica e materiais para teste de usabilidade.",
        ],
        "evidence": "npm run test:a11y, testes de teclado, readability test e docs/ethics.",
        "report_link": "Pode alimentar Ethics, Accessibility, Evaluation Plan e Results.",
    },
    {
        "week": "Semana 6 / refinamento atual",
        "focus": "Fontes CPAG, UI moderna e prontidao para demonstracao",
        "features": [
            "CPAG Welfare Rights definido como fonte principal exibida no rule trace.",
            "GOV.UK mantido como fonte oficial complementar.",
            "Remake visual Blue Decision Lab com linguagem tecnologica azul.",
            "Cards de resultados polidos com badge curto, superficie mais limpa e acoes menos pesadas.",
            "Export JSON anonimo para apoiar teste de usabilidade.",
            "Deploy de producao atualizado em https://verissimo-tawny.vercel.app.",
        ],
        "evidence": "119 testes passando, tsc, lint, coverage, build, test:a11y e smoke de producao.",
        "report_link": "Pode alimentar Final Artefact, Evaluation Readiness e Demonstration.",
    },
    {
        "week": "Semana 7 / proximo passo",
        "focus": "Teste com participantes e coleta de feedback",
        "features": [
            "Submeter e aguardar aprovacao etica antes de envolver participantes reais.",
            "Executar teste com 5 a 8 participantes usando cenarios ficticios.",
            "Coletar export JSON anonimo e respostas do questionario pos-teste.",
            "Analisar compreensao, confianca, acessibilidade percebida e pontos de confusao.",
        ],
        "evidence": "Consent form, participant information sheet, script, tasks e questionnaire.",
        "report_link": "Vai alimentar Methodology, Evaluation, Results e Further Work.",
    },
]


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_executive_summary(doc)
    add_weekly_feature_plan(doc)
    add_architecture_summary(doc)
    add_data_sources_summary(doc)
    add_testing_summary(doc)
    add_ethics_summary(doc)
    add_report_writing_map(doc)
    add_risks_and_next_steps(doc)
    add_appendix_evidence(doc)
    add_footer(doc)
    doc.save(OUTPUT)
    structural_audit(OUTPUT)
    print(OUTPUT)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.start_type = WD_SECTION_START.NEW_PAGE

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.bold = True
        style.paragraph_format.keep_with_next = True

    h1 = doc.styles["Heading 1"]
    h1.font.size = Pt(16)
    h1.font.color.rgb = BLUE
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)

    h2 = doc.styles["Heading 2"]
    h2.font.size = Pt(13)
    h2.font.color.rgb = BLUE
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    h3 = doc.styles["Heading 3"]
    h3.font.size = Pt(12)
    h3.font.color.rgb = DARK_BLUE
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("VERISSIMO")
    set_run_font(r, size=12, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Relatorio de apoio semanal")
    set_run_font(r, size=24, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Aplicacao web para apoio a compreensao de beneficios sociais do Reino Unido")
    set_run_font(r, size=13, color=MUTED)

    rows = [
        ("Projeto", "VERISSIMO"),
        ("Aluna", "Viviane Verissimo dos Santos"),
        ("Curso", "BSc (Hons) Computing - Southampton Solent University"),
        ("Data", "20/07/2026"),
        ("Link publico", "https://verissimo-tawny.vercel.app"),
        ("Objetivo do documento", "Apoiar apresentacoes semanais e servir como base para a escrita do report academico."),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    set_table_widths(table, [1800, 7560])
    for idx, (label, value) in enumerate(rows):
        table.cell(idx, 0).text = label
        table.cell(idx, 1).text = value
        shade_cell(table.cell(idx, 0), LIGHT_GRAY)
        for cell in table.rows[idx].cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        bold_first_run(table.cell(idx, 0))

    doc.add_paragraph()
    add_callout(
        doc,
        "Mensagem central",
        "O VERISSIMO separa decisao e explicacao: um motor de regras deterministico calcula os resultados, enquanto a IA e usada apenas para reescrever explicacoes em linguagem mais simples. Essa separacao torna o prototipo mais auditavel, seguro e adequado para uma avaliacao academica.",
        LIGHT_BLUE,
    )


def add_executive_summary(doc: Document) -> None:
    add_heading(doc, "1. Resumo executivo", 1)
    add_paragraph(
        doc,
        "O VERISSIMO e um prototipo web que ajuda usuarios a revisar possiveis beneficios sociais do Reino Unido. A aplicacao guia a pessoa por perguntas simples, avalia as respostas contra regras declarativas e apresenta resultados com fontes verificaveis, evidencias recomendadas e explicacoes simplificadas.",
    )
    add_paragraph(
        doc,
        "A versao atual cobre Attendance Allowance, Universal Credit e Personal Independence Payment. O projeto prioriza auditabilidade, privacidade, acessibilidade e clareza. Ele nao substitui GOV.UK, CPAG, DWP, welfare advisers ou decisoes oficiais.",
    )
    add_bullets(
        doc,
        [
            "Artefato principal: aplicacao Next.js publicada na Vercel.",
            "Metodo de decisao: motor de regras deterministico baseado em YAML.",
            "Fonte principal exibida no trace: CPAG Welfare Rights.",
            "Fonte complementar: GOV.UK.",
            "IA: Gemini opcional, limitada a explicacao; nao decide elegibilidade.",
            "Estado atual: 119 testes locais passando e deploy publico ativo.",
        ],
    )


def add_weekly_feature_plan(doc: Document) -> None:
    add_heading(doc, "2. Features por semana", 1)
    add_paragraph(
        doc,
        "Esta secao organiza o progresso em blocos semanais para facilitar reunioes, apresentacoes e escrita incremental do report. Cada semana tem um foco, as features entregues, a evidencia tecnica e a parte da dissertacao que ela pode alimentar.",
    )
    for item in WEEKLY_FEATURES:
        add_heading(doc, f"{item['week']} - {item['focus']}", 2)
        add_bullets(doc, item["features"])
        add_label_value(doc, "Evidencia", item["evidence"])
        add_label_value(doc, "Uso no report", item["report_link"])

    add_heading(doc, "Resumo tabular semanal", 2)
    rows = [["Semana", "Foco", "Entregas principais", "Evidencia"]]
    for item in WEEKLY_FEATURES:
        rows.append(
            [
                item["week"],
                item["focus"],
                "; ".join(item["features"][:3]),
                item["evidence"],
            ]
        )
    add_table(doc, rows, [1150, 1900, 4010, 2300])


def add_architecture_summary(doc: Document) -> None:
    add_heading(doc, "3. Arquitetura da aplicacao", 1)
    add_paragraph(
        doc,
        "A arquitetura foi desenhada para manter as decisoes de elegibilidade fora do LLM. O frontend apresenta o questionario e os resultados. A API interna recebe respostas e chama o motor de regras no servidor. Os arquivos YAML armazenam os criterios declarativos por beneficio.",
    )
    rows = [
        ["Camada", "Responsabilidade", "Arquivos principais"],
        ["Frontend", "Questionario guiado, cards de resultado, relatorio imprimivel e export anonimo.", "src/components/assessment/AssessmentApp.tsx; src/app/globals.css"],
        ["API", "Rotas internas para avaliacao e explicacao.", "src/app/api/assess/route.ts; src/app/api/explain/route.ts"],
        ["Motor de regras", "Carregar YAML, validar schema, avaliar criterios, derivar outcome e evidencias.", "src/lib/rules/engine.ts; load.ts; schema.ts; outcome.ts; operators.ts"],
        ["Regras", "Criterios declarativos por beneficio, fontes e evidencias.", "rules/attendance-allowance.yaml; universal-credit.yaml; personal-independence-payment.yaml"],
        ["Explicacao", "Template deterministico, prompt seguro, validador e adaptador Gemini.", "src/lib/explanations/template.ts; llm.ts; gemini.ts"],
        ["Testes", "Unitarios, rotas, UI, acessibilidade, teclado e legibilidade.", "tests/rules; tests/intake; tests/explanations; tests/e2e"],
    ]
    add_table(doc, rows, [1450, 3600, 4310])

    add_callout(
        doc,
        "Decisao arquitetural importante",
        "O LLM nao calcula elegibilidade, nao cria criterio e nao escolhe beneficio. Ele recebe um resultado ja calculado e tenta explicar em texto mais simples. Se a resposta externa for insegura ou indisponivel, a aplicacao usa fallback deterministico.",
        LIGHT_YELLOW,
    )


def add_data_sources_summary(doc: Document) -> None:
    add_heading(doc, "4. Fontes de dados e rastreabilidade", 1)
    add_paragraph(
        doc,
        "A fonte principal exibida para regras de welfare rights e CPAG. GOV.UK continua como fonte oficial complementar. A aplicacao nao faz scraping ao vivo: as regras ficam explicitamente modeladas em YAML para que o comportamento seja reproduzivel e testavel.",
    )
    rows = [
        ["Beneficio", "Fonte CPAG principal", "Fonte GOV.UK complementar"],
        ["Universal Credit", "https://cpag.org.uk/welfare-rights/key-topics/universal-credit/universal-credit-basics", "https://www.gov.uk/universal-credit/eligibility"],
        ["Personal Independence Payment", "https://cpag.org.uk/welfare-rights/key-topics/personal-independence-payment", "https://www.gov.uk/pip/eligibility"],
        ["Attendance Allowance", "https://cpag.org.uk/welfare-rights", "https://www.gov.uk/attendance-allowance/eligibility"],
    ]
    add_table(doc, rows, [1800, 4180, 3380])
    add_bullets(
        doc,
        [
            "Cada resultado carrega rulesVersion para rastrear a versao exata do YAML usado.",
            "Cada criterio tem explicacao e URL de fonte.",
            "Fontes complementares aparecem separadas do trace CPAG.",
            "A decisao de nao usar scraping ao vivo reduz instabilidade durante avaliacao e demonstracao.",
        ],
    )


def add_testing_summary(doc: Document) -> None:
    add_heading(doc, "5. Evidencias de qualidade", 1)
    add_paragraph(
        doc,
        "O projeto possui verificacoes automatizadas que podem ser citadas no report como evidencia de engenharia, confiabilidade e acessibilidade. A verificacao mais recente, apos o polimento dos cards de resultado, passou nos seguintes gates.",
    )
    rows = [
        ["Verificacao", "Resultado recente", "O que comprova"],
        ["npm test", "14 arquivos / 119 testes passando", "Regras, rotas, UI, explicacoes, helpers e legibilidade."],
        ["npx tsc --noEmit", "Passou", "Consistencia de tipos TypeScript."],
        ["npm run lint", "Passou", "Padrao estatico de codigo."],
        ["npm run test:coverage", "97.43% statements; 94.2% branches", "Cobertura acima do gate definido."],
        ["npm run build", "Passou", "Build de producao Next.js."],
        ["npm run test:a11y", "3 Playwright tests passando", "Axe/WCAG automatizado e fluxo por teclado."],
        ["Smoke de producao", "HTTP 200 e API retornando CPAG + GOV.UK", "Deploy publico funcional apos alteracoes."],
    ]
    add_table(doc, rows, [2300, 2700, 4360])


def add_ethics_summary(doc: Document) -> None:
    add_heading(doc, "6. Etica, privacidade e avaliacao com usuarios", 1)
    add_paragraph(
        doc,
        "O projeto foi estruturado para reduzir risco antes de qualquer teste com participantes. Nao ha login, banco de dados de usuarios ou armazenamento persistente de dados pessoais. Os testes devem usar cenarios ficticios ate haver aprovacao etica.",
    )
    add_bullets(
        doc,
        [
            "Participantes nao devem inserir dados reais de saude, renda, imigracao ou beneficios.",
            "O export JSON e anonimo e serve para avaliacao de usabilidade, nao para perfilamento.",
            "O site informa que nao oferece aconselhamento oficial.",
            "O pacote docs/ethics contem consent form, participant information sheet, screening, task list, script e post-test questionnaire.",
        ],
    )
    add_callout(
        doc,
        "Ponto para reuniao semanal",
        "Antes de executar usability testing real, o proximo marco e revisar e submeter o pacote de etica no portal da universidade.",
        LIGHT_RED,
    )


def add_report_writing_map(doc: Document) -> None:
    add_heading(doc, "7. Como transformar isso no report academico", 1)
    rows = [
        ["Secao do report", "Conteudo que pode ser usado", "Evidencias locais"],
        ["Introduction / Problem", "Complexidade de beneficios, risco de IA tomar decisoes e necessidade de explicabilidade.", "Resumo executivo e limitacoes do prototipo."],
        ["Requirements", "Tres beneficios, fontes rastreaveis, acessibilidade, privacidade, explicacao simples e deploy publico.", "Specs em docs/specs."],
        ["Methodology", "Desenvolvimento incremental, TDD, verificacao continua, UI reviews e plano de usability testing.", "WORKLOG.md e testes automatizados."],
        ["Design and Implementation", "Arquitetura Next.js, API interna, motor YAML, CPAG/GOV.UK, LLM boundary e UI Blue Decision Lab.", "src/, rules/, docs/plans e screenshots."],
        ["Testing", "Vitest, coverage, TypeScript, ESLint, Playwright, axe-core, keyboard-only flow e readability audit.", "tests/ e coverage/."],
        ["Evaluation", "Plano de 5 a 8 participantes com cenarios ficticios e export JSON anonimo.", "docs/ethics e export study JSON."],
        ["Conclusion / Further Work", "Mais beneficios, painel de regras, revisao com adviser, atualizacao de fontes e avaliacao longitudinal.", "Riscos e proximos passos deste documento."],
    ]
    add_table(doc, rows, [1800, 4470, 3090])


def add_risks_and_next_steps(doc: Document) -> None:
    add_heading(doc, "8. Riscos, limitacoes e proximos passos", 1)
    rows = [
        ["Nivel", "Risco ou limitacao", "Mitigacao proposta"],
        ["Alto", "Regras de beneficios podem mudar e o prototipo nao deve dar aconselhamento oficial.", "Registrar last_verified, mostrar fontes e recomendar CPAG/GOV.UK/adviser."],
        ["Alto", "Teste com usuarios reais sem aprovacao etica.", "Usar apenas cenarios ficticios ate aprovacao formal."],
        ["Medio", "LLM pode produzir explicacao inadequada.", "Validacao de resposta, allowlist de URLs e fallback deterministico."],
        ["Medio", "Cobertura de apenas tres beneficios.", "Apresentar como prototipo focado e listar expansao como further work."],
        ["Baixo", "UI pode precisar de ajustes apos feedback de participantes.", "Usar export JSON, questionario pos-teste e iteracoes visuais documentadas."],
    ]
    add_table(doc, rows, [900, 4300, 4160])

    add_heading(doc, "Proximos passos recomendados", 2)
    add_numbered(
        doc,
        [
            "Revisar este relatorio semanal com supervisor ou stakeholder.",
            "Finalizar submissao etica.",
            "Preparar roteiro de demonstracao usando o site publico.",
            "Executar usability testing com 5 a 8 participantes apos aprovacao.",
            "Transformar as secoes deste documento em capitulos do report final.",
        ],
    )


def add_appendix_evidence(doc: Document) -> None:
    add_heading(doc, "9. Evidencias e artefatos principais", 1)
    add_bullets(
        doc,
        [
            "Aplicacao publica: https://verissimo-tawny.vercel.app",
            "Regras: rules/attendance-allowance.yaml, rules/universal-credit.yaml, rules/personal-independence-payment.yaml",
            "Motor: src/lib/rules/",
            "UI: src/components/assessment/AssessmentApp.tsx e src/app/globals.css",
            "LLM: src/lib/explanations/",
            "Testes: tests/",
            "Etica: docs/ethics/",
            "Evidencia visual recente: test-results/result-card-polish-desktop.png",
            "Ultimo deploy citado: dpl_EsNxnGJ5kKT31J4hwStxprmnyJdt",
        ],
    )


def add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)


def add_label_value(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f"{label}: ")
    r.bold = True
    p.add_run(value)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_callout(doc: Document, label: str, text: str, fill: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p.add_run(text)
    doc.add_paragraph()


def add_table(doc: Document, rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    set_table_widths(table, widths)
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = value
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=10 if row_idx else 10.5, color=WHITE if row_idx == 0 else INK, bold=row_idx == 0)
            if row_idx == 0:
                shade_cell(cell, "2E74B5")
            elif row_idx % 2 == 0:
                shade_cell(cell, "F8FAFC")
    doc.add_paragraph()


def set_run_font(run, size=None, color=None, bold=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def bold_first_run(cell) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True


def set_table_widths(table, widths: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("VERISSIMO - weekly support report")
    set_run_font(r, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Generated for weekly presentation support")
    set_run_font(r, size=9, color=MUTED)


def structural_audit(path: Path) -> None:
    with ZipFile(path) as docx:
        document_xml = docx.read("word/document.xml").decode("utf-8")
        styles_xml = docx.read("word/styles.xml").decode("utf-8")
    required = [
        "Relatorio de apoio semanal",
        "Features por semana",
        "Semana 1",
        "Semana 6 / refinamento atual",
        "CPAG Welfare Rights",
        "119 testes locais passando",
        "Como transformar isso no report academico",
    ]
    missing = [text for text in required if text not in document_xml]
    if missing:
        raise RuntimeError(f"Missing required document text: {missing}")
    forbidden = ["TODO", "TBD", "PLACEHOLDER"]
    found = [text for text in forbidden if text in document_xml]
    if found:
        raise RuntimeError(f"Forbidden placeholders found: {found}")
    if "Heading1" not in styles_xml and "Heading 1" not in styles_xml:
        raise RuntimeError("Heading styles not present in styles.xml")


if __name__ == "__main__":
    main()
